"""
run_pipeline.py  (Phase 2 rewrite)

Runs the full sentiment pipeline on a schedule instead of needing a
manual re-run every time - Phase 2 correction #1 ("Replace manual
Excel updates with a scheduled pull") and the target architecture
("each pipeline runs itself on an interval... writes its latest
output to a shared store").

WHAT'S NEW VS PHASE 1
----------------------
- Default interval is 10 minutes (was a 30-min --watch flag with no
  default schedule) - matches "e.g. every 10 minutes" in the roadmap.
- Every fetch funnels through sentiment_engine.process_item(), which
  dedups (URL + near-duplicate headline) BEFORE scoring and writes
  straight into sentiment_store.py's SQLite tables - not into memory
  waiting to be dumped to Excel.
- Every run also: (a) snapshots the current per-asset average sentiment
  into sentiment_trend (this is what powers "now / 1h ago / 6h ago /
  24h ago"), and (b) checks that new snapshot against the previous one
  and fires an alert row if any asset's sentiment swung sharply
  (default: |delta| >= 0.35) - Phase 2 correction #7.
- Excel is now export-only (export_snapshot_to_excel), matching
  correction #8: "Excel becomes an export option only, never the thing
  the dashboard reads from." Run with --export-excel to also produce a
  static workbook alongside the live database, e.g. for the report.

Usage:
    python run_pipeline.py                  # one snapshot run, then exit
    python run_pipeline.py --watch           # repeats every 10 min until Ctrl+C
    python run_pipeline.py --watch --interval 5
    python run_pipeline.py --export-excel    # also write an Excel snapshot
"""

import argparse
import time
from datetime import datetime

import sentiment_engine as engine
import sentiment_store as store

# ---------------------------------------------------------------------
# CONFIG - fill these in with your own credentials / targets
# ---------------------------------------------------------------------
CONFIG = {
    "telegram": {
        "enabled": False,
        "api_id": None,
        "api_hash": None,
        "channels": ["@cointelegraph", "@CoinDeskGlobal", "@WatcherGuru"],
        "limit_per_channel": 30,
    },
    "youtube": {
        "enabled": True,
        "api_key": "AIzaSyDCfBVyv4TDClRfkJIBTrfj0YdAz8b4kGw",
        "mode": "keyword",
        "keywords": ["gold price today", "crypto crash", "stock market fraud india"],
        "max_videos_per_keyword": 5,
        "include_comments": True,
        "comments_per_video": 30,
        "video_ids": ["dQw4w9WgXcQ"],
        "max_comments": 30,
    },
    "news": {
        "enabled": True,
        "mode": "keyword",
        "keywords": ["gold price", "cryptocurrency crash", "stock market fraud india", "sensex today"],
        "region": "IN",
        "only_relevant": True,
        "feeds": engine.DEFAULT_NEWS_FEEDS,
    },
    "images": {
        "enabled": True,
        "paths": [],
    },
}

# Assets tracked for trend snapshots + swing alerts every run.
TRACKED_ASSETS = ["Gold", "Crypto", "Market", "Fraud"]

# Minimum change in average sentiment (on the -1..+1 scale) between this
# run and the previous snapshot to fire a swing alert.
SWING_ALERT_THRESHOLD = 0.35


def fetch_new_items(quiet: bool = False) -> list:
    """Pulls from every enabled source. Each adapter already dedups and
    writes to the shared store internally (see sentiment_engine.py), so
    what comes back here is only what was actually NEW this run."""
    new_results = []

    if CONFIG["telegram"]["enabled"]:
        for channel in CONFIG["telegram"]["channels"]:
            if not quiet:
                print(f"Fetching Telegram: {channel}")
            new_results += engine.analyze_telegram_channel(
                channel, CONFIG["telegram"]["api_id"], CONFIG["telegram"]["api_hash"],
                limit=CONFIG["telegram"]["limit_per_channel"],
            )

    if CONFIG["youtube"]["enabled"]:
        yt = CONFIG["youtube"]
        if yt["mode"] == "keyword":
            for kw in yt["keywords"]:
                if not quiet:
                    print(f'Searching YouTube: "{kw}"')
                new_results += engine.analyze_youtube_by_keyword(
                    kw, yt["api_key"], max_videos=yt["max_videos_per_keyword"],
                    include_comments=yt["include_comments"], comments_per_video=yt["comments_per_video"],
                )
        else:
            for vid in yt["video_ids"]:
                if not quiet:
                    print(f"Fetching YouTube comments: {vid}")
                new_results += engine.analyze_youtube_comments(vid, yt["api_key"], yt["max_comments"])

    if CONFIG["news"]["enabled"]:
        news = CONFIG["news"]
        if news["mode"] == "keyword":
            if not quiet:
                print(f"Searching news for: {', '.join(news['keywords'])}")
            new_results += engine.analyze_news_by_keyword(
                news["keywords"], region=news["region"], only_relevant=news["only_relevant"],
            )
        else:
            if not quiet:
                print("Fetching news RSS feeds...")
            new_results += engine.analyze_news_feed(news["feeds"])

    if CONFIG["images"]["enabled"]:
        for path in CONFIG["images"]["paths"]:
            if not quiet:
                print(f"OCR-ing image: {path}")
            new_results += engine.analyze_image(path)

    return new_results


def snapshot_trend_and_check_alerts(quiet: bool = False):
    """Writes this run's per-asset average sentiment into sentiment_trend,
    and compares it against the previous snapshot for that asset - firing
    an alert row (and printing it) on a sharp swing. This is what makes
    'Bitcoin sentiment +0.45 -> -0.32' a real, queryable event instead of
    something you'd only notice by eyeballing two Excel files."""
    summary = store.get_asset_summary()

    for asset in TRACKED_ASSETS:
        stats = summary.get(asset)
        if not stats or stats["count"] == 0:
            continue

        avg_sentiment = round(stats["avg_sentiment"], 3)
        previous = store.get_latest_trend_value(asset)

        store.insert_trend_snapshot(
            asset=asset, avg_sentiment=avg_sentiment,
            item_count=stats["count"], high_risk_count=stats["high_risk_count"],
        )

        if previous is not None:
            delta = round(avg_sentiment - previous["avg_sentiment"], 3)
            if abs(delta) >= SWING_ALERT_THRESHOLD:
                direction = "up" if delta > 0 else "down"
                severity = "High" if abs(delta) >= 0.6 else "Medium"
                message = (f"{asset} sentiment swung {direction} by {delta:+.2f} "
                           f"({previous['avg_sentiment']:+.2f} -> {avg_sentiment:+.2f})")
                store.insert_alert(
                    asset=asset, alert_type="sentiment_swing", message=message,
                    old_value=previous["avg_sentiment"], new_value=avg_sentiment,
                    delta=delta, severity=severity,
                )
                if not quiet:
                    print(f"  ALERT [{severity}] {message}")


def export_snapshot_to_excel(out_path: str = "sentiment_report.xlsx", limit: int = 500):
    """Export-only, per Phase 2 correction #8 - this is NOT what the
    dashboard reads from; it exists for attaching a static snapshot to
    the report/submission. Pulls straight from the shared store so the
    export always matches what's live."""
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Latest Sentiment"
    headers = ["Fetched At", "Source Type", "Source", "Asset", "Sentiment (FinBERT)",
               "Score", "Confidence", "Risk Level", "Risk Score", "Headline / Text"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    fills = {
        "High": PatternFill(start_color="F8CBCB", end_color="F8CBCB", fill_type="solid"),
        "Medium": PatternFill(start_color="FCEBB6", end_color="FCEBB6", fill_type="solid"),
        "Low": PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"),
    }

    for row in store.get_latest_news(limit=limit):
        ws.append([
            row["fetched_at"], row["source_type"], row["source_ref"], row["asset"],
            row["final_label"], row["final_score"], row["finbert_confidence"],
            row["risk_level"], row["risk_score"], (row["headline"] or row["raw_text"] or "")[:500],
        ])
        ws.cell(row=ws.max_row, column=8).fill = fills.get(row["risk_level"])

    for col, width in zip("ABCDEFGHIJ", [20, 14, 22, 10, 16, 8, 10, 10, 10, 60]):
        ws.column_dimensions[col].width = width

    ws2 = wb.create_sheet("Asset Summary")
    ws2.append(["Asset", "Item Count", "Avg Sentiment", "High-Risk Items"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for asset, s in store.get_asset_summary().items():
        ws2.append([asset, s["count"], round(s["avg_sentiment"], 3), s["high_risk_count"]])
    for col, width in zip("ABCD", [14, 12, 14, 16]):
        ws2.column_dimensions[col].width = width

    ws3 = wb.create_sheet("Recent Alerts")
    ws3.append(["Created At", "Asset", "Severity", "Message"])
    for cell in ws3[1]:
        cell.font = Font(bold=True)
    for a in store.get_recent_alerts(limit=50):
        ws3.append([a["created_at"], a["asset"], a["severity"], a["message"]])
    for col, width in zip("ABCD", [20, 12, 10, 70]):
        ws3.column_dimensions[col].width = width

    wb.save(out_path)
    return out_path


def run(quiet: bool = False, export_excel: bool = False):
    """One full cycle: fetch new items (dedup + score + store happens
    inside fetch_new_items), snapshot the trend, check for swing alerts,
    and optionally export a static Excel copy."""
    store.init_db()
    new_results = fetch_new_items(quiet=quiet)

    if not quiet:
        print(f"\nNew items scored this run: {len(new_results)}")

    snapshot_trend_and_check_alerts(quiet=quiet)

    if not quiet:
        summary = store.get_asset_summary()
        for asset, s in summary.items():
            print(f"  {asset:<10} count={s['count']:<4} avg_sentiment={round(s['avg_sentiment'], 3):<7} "
                  f"high_risk={s['high_risk_count']}")

    if export_excel:
        path = export_snapshot_to_excel()
        if not quiet:
            print(f"\nExcel snapshot written to: {path}")

    return new_results


def run_scheduled(interval_minutes: int = 10, iterations: int = None, export_excel: bool = False):
    """Re-runs run() on a fixed schedule so the shared store - and
    therefore the dashboard - updates on its own, with no manual re-run.
    For a real deployment, a cron job / systemd timer / Windows Task
    Scheduler is more robust than leaving a terminal open, but this loop
    is enough for a live demo/viva."""
    count = 0
    print(f"Starting scheduled run every {interval_minutes} minute(s). Press Ctrl+C to stop.")
    try:
        while iterations is None or count < iterations:
            count += 1
            print(f"\n=== Run {count} at {datetime.now().isoformat(timespec='seconds')} ===")
            run(export_excel=export_excel)
            if iterations is None or count < iterations:
                time.sleep(interval_minutes * 60)
    except KeyboardInterrupt:
        print("\nStopped.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--watch", action="store_true", help="run on a repeating schedule")
    parser.add_argument("--interval", type=int, default=10, help="minutes between runs when --watch is set")
    parser.add_argument("--export-excel", action="store_true", help="also write a static Excel snapshot")
    args = parser.parse_args()

    if args.watch:
        run_scheduled(interval_minutes=args.interval, export_excel=args.export_excel)
    else:
        run(export_excel=args.export_excel)
